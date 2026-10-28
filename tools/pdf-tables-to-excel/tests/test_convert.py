"""PyTest: builds a 3-page sample PDF with headless Edge, converts it, checks cells.
Run from the project folder:  py -m pytest -q
"""
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pdf_tables_to_excel import convert, number_or_text  # noqa: E402

HERE = Path(__file__).resolve().parent
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")


@pytest.fixture(scope="module")
def sample_pdf(tmp_path_factory):
    if not EDGE.exists():
        pytest.skip("Edge not installed; cannot render the sample PDF")
    out = tmp_path_factory.mktemp("pdf") / "sample.pdf"
    subprocess.run([str(EDGE), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={out}", (HERE / "make_sample_pdf.html").as_uri()],
                   check=False, capture_output=True, timeout=60)
    assert out.exists() and out.stat().st_size > 1000
    return out


def test_one_sheet_per_page_and_values(sample_pdf, tmp_path):
    s = convert(sample_pdf, tmp_path, force_text=False)
    assert s["pages"] == 3 and s["empty_pages"] == []
    wb = load_workbook(s["output"])
    assert wb.sheetnames == ["Page 1", "Page 2", "Page 3"]

    p1 = wb["Page 1"]
    assert p1["A1"].value == "Item" and p1["D1"].value == "Total"
    assert p1["A2"].value == "Solar panel 550 W"
    assert p1["B2"].value == 10                      # numeric, not text
    assert p1["D3"].value == 168000                  # thousands separator removed
    assert p1["D4"].value == pytest.approx(21750.50)
    # second table stacked after a blank row
    assert p1["A5"].value is None and p1["A6"].value == "Code" and p1["B8"].value == "Pending"

    p2 = wb["Page 2"]
    assert p2["A1"].value == "Load" and p2["B3"].value == 100

    p3 = wb["Page 3"]
    assert 3 in s["text_mode_pages"]
    texts = [c.value for row in p3.iter_rows() for c in row if c.value]
    assert any("paragraphs" in str(t) for t in texts)


def test_number_parsing_keeps_codes_as_text():
    assert number_or_text("1,250") == 1250
    assert number_or_text("21750.50") == 21750.5
    assert number_or_text("0") == 0
    assert number_or_text("007") == "007"            # leading zeros are codes, keep text
    assert number_or_text("A-001") == "A-001"
    assert number_or_text("-15") == -15
